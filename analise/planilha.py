"""
analise/planilha.py — Reúne as tabelas do estudo de caso numa planilha.

Gera `media/exports/estudo_caso.xlsx` com uma aba por tabela do artigo, a
partir dos mesmos CSVs que os demais scripts desta pasta. Os criadores saem
como C1…C10.

A planilha não é versionada: é gerada sob demanda a partir de `dados/`, de
modo que não existe a possibilidade de o arquivo binário e os CSVs contarem
histórias diferentes. Quem quiser as tabelas em texto roda os scripts
individuais, que imprimem o mesmo conteúdo no terminal.

Uso:
    python analise/planilha.py
"""

import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from comum import METRICAS, RAIZ, carregar, ler_csv
from dtw import sequencias
from tabela2 import testar
from tabela3 import NEUTRA, distribuicao
from tabela4 import agregar

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SAIDA = os.path.join(RAIZ, "media", "exports", "estudo_caso.xlsx")

CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO = PatternFill("solid", fgColor="1F3864")


def _escrever(ws, colunas: list, linhas: list) -> None:
    ws.append(colunas)
    for celula in ws[1]:
        celula.font = CABECALHO
        celula.fill = FUNDO
        celula.alignment = Alignment(horizontal="center")
    for linha in linhas:
        ws.append(linha)
    for i, nome in enumerate(colunas, start=1):
        largura = max([len(str(nome))] + [len(str(l[i - 1])) for l in linhas] or [0])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(largura + 3, 48)
    ws.freeze_panes = "A2"


def aba_tabela2(wb, dados: dict) -> None:
    ws = wb.create_sheet("Tabela 2 — grupos")
    linhas = [[d["metrica"], round(d["mediana_a"], 3), round(d["mediana_b"], 3),
               int(d["u"]), d["p"], round(d["rb"], 3), round(d["q"], 6),
               "sobrevive" if d["q"] < 0.05 else "cai"]
              for d in testar(dados)]
    _escrever(ws, ["Métrica", "Mediana A", "Mediana B", "U", "p",
                   "r (rank-biserial)", "q (Benjamini-Hochberg)", "Veredito"], linhas)


def aba_tabela3(wb, dados: dict) -> None:
    ws = wb.create_sheet("Tabela 3 — emoções")
    contagem = distribuicao({v: d["grupo"] for v, d in dados.items()})
    linhas = []
    for grupo in ("A", "B"):
        c = contagem[grupo]
        total = sum(c.values())
        nao_neutros = total - c.get(NEUTRA, 0)
        for emocao, n in c.most_common():
            linhas.append([
                grupo, emocao, n, round(n / total * 100, 1),
                "—" if emocao == NEUTRA else round(n / nao_neutros * 100, 1),
            ])
    _escrever(ws, ["Grupo", "Emoção", "Segmentos", "% do total",
                   "% dos não-neutros"], linhas)


def aba_tabela4(wb, dados: dict) -> None:
    ws = wb.create_sheet("Tabela 4 — audiência")
    linhas = [[d["criador"], d["grupo"], d["total"], round(d["sentimento"], 3),
               round(d["pos"] * 100, 1), round(d["neg"] * 100, 1),
               d["emocao"], d["palavras"]]
              for d in agregar(dados)]
    _escrever(ws, ["Criador", "Grupo", "Comentários analisados",
                   "Sentimento médio", "% positivos", "% negativos",
                   "Emoção dominante", "Top 5 palavras"], linhas)


def aba_fingerprint(wb, dados: dict) -> None:
    """Uma linha por vídeo: o fingerprint como entra nas análises."""
    ws = wb.create_sheet("Fingerprint por vídeo")
    colunas = ["Vídeo", "Grupo", "Criador"] + list(METRICAS)
    linhas = []
    for vid in sorted(dados, key=lambda v: (dados[v]["grupo"], dados[v]["criador"], v)):
        d = dados[vid]
        linhas.append([vid, d["grupo"], d["criador"]]
                      + [round(d[m], 4) if d[m] is not None else "" for m in METRICAS])
    _escrever(ws, colunas, linhas)


def aba_dtw(wb, dados: dict) -> None:
    ws = wb.create_sheet("Sequências emocionais")
    grupo_de = {v: d["grupo"] for v, d in dados.items()}
    seqs = sequencias(grupo_de)
    linhas = []
    for grupo in ("A", "B"):
        for vid in sorted(seqs[grupo]):
            codigos = seqs[grupo][vid]
            linhas.append([vid, grupo, dados[vid]["criador"], len(codigos),
                           " ".join(str(c) for c in codigos)])
    _escrever(ws, ["Vídeo", "Grupo", "Criador", "Segmentos",
                   "Sequência (codificação ordinal)"], linhas)


def main() -> int:
    dados = carregar()
    wb = Workbook()
    wb.remove(wb.active)

    aba_tabela2(wb, dados)
    aba_tabela3(wb, dados)
    aba_tabela4(wb, dados)
    aba_fingerprint(wb, dados)
    aba_dtw(wb, dados)

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    wb.save(SAIDA)

    print(f"Planilha gerada: {SAIDA}")
    for ws in wb.worksheets:
        print(f"   {ws.title:32} {ws.max_row - 1:>5} linhas")
    print()
    print(f"   {len(ler_csv('corpus_estudo'))} vídeos · criadores como C1…C10")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
